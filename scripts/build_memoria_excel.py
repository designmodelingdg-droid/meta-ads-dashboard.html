"""Hoja de resumen de verificaciones que acompana a la plantilla de memoria.

La hoja NO calcula la ingenieria: calcula el ratio y decide CUMPLE / NO CUMPLE.
Eso es a proposito. En el momento en que la hoja propusiera una seccion o un
valor de diseno, dejaria de ser una plantilla documental y pasaria a ser una
herramienta de calculo — y entonces tendria que responder por el numero.

    python3 scripts/build_memoria_excel.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

NAVY, ORANGE, GREY, CREMA = "FF0E2438", "FFC96A1C", "FF5A6B7B", "FFFAFAF7"
OK_BG, OK_TX = "FFE2F0E9", "FF146B48"
NO_BG, NO_TX = "FFFDF0EE", "FF8F2F1E"

wb = Workbook()
thin = Side(style="thin", color="FFD9DEE3")
box  = Border(left=thin, right=thin, top=thin, bottom=thin)

def hoja(ws, titulo, sub, cabeceras, anchos, filas, formula_desde=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "DESIGN MODELING ACADEMY"
    ws["A1"].font = Font(name="Overpass", bold=True, size=9, color=ORANGE)
    ws["A2"] = titulo
    ws["A2"].font = Font(name="Overpass", bold=True, size=16, color=NAVY)
    ws["A3"] = sub
    ws["A3"].font = Font(name="Nunito", size=10, color=GREY, italic=True)
    ws.row_dimensions[2].height = 22

    r = 5
    for i, c in enumerate(cabeceras, 1):
        cel = ws.cell(row=r, column=i, value=c)
        cel.font = Font(name="Overpass", bold=True, size=9, color="FFFFFFFF")
        cel.fill = PatternFill("solid", fgColor=NAVY)
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        cel.border = box
    ws.row_dimensions[r].height = 30
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a

    for j, fila in enumerate(filas):
        rr = r + 1 + j
        for i, v in enumerate(fila, 1):
            cel = ws.cell(row=rr, column=i, value=v)
            cel.font = Font(name="Nunito", size=10)
            cel.border = box
            cel.alignment = Alignment(vertical="center", wrap_text=(i == 1))
            if isinstance(v, str) and v.startswith("["):
                cel.font = Font(name="Nunito", size=10, color="FFB9C2CB", italic=True)
    return r + 1, r + len(filas)

# ── ELU ──────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "ELU · Resistencia"
vacias = [["[ elemento ]", "[ tipo ]", None, None, None, None, "[ norma y artículo ]"] for _ in range(14)]
ini, fin = hoja(ws, "Resumen de verificaciones · Estado límite último",
    "Rellena solicitación, resistencia y límite. El ratio y el veredicto se calculan solos.",
    ["Elemento", "Verificación", "Solicitación\n(Sd)", "Resistencia\n(Rd)", "Ratio\nSd/Rd", "Veredicto", "Referencia normativa"],
    [22, 24, 13, 13, 10, 13, 30], vacias)

for r in range(ini, fin + 1):
    ws.cell(row=r, column=5, value=f'=IF(N(D{r})=0,"",C{r}/D{r})').number_format = "0.00"
    # El limite es 1,00: por encima, la solicitacion supera a la resistencia.
    ws.cell(row=r, column=6, value=f'=IF(E{r}="","",IF(E{r}<=1,"CUMPLE","NO CUMPLE"))')
    ws.cell(row=r, column=6).font = Font(name="Overpass", bold=True, size=9)
    ws.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="center")

rng = f"F{ini}:F{fin}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CUMPLE"'],
    fill=PatternFill("solid", fgColor=OK_BG), font=Font(name="Overpass", bold=True, size=9, color=OK_TX)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NO CUMPLE"'],
    fill=PatternFill("solid", fgColor=NO_BG), font=Font(name="Overpass", bold=True, size=9, color=NO_TX)))
ws.conditional_formatting.add(f"E{ini}:E{fin}", CellIsRule(operator="greaterThan", formula=["1"],
    font=Font(name="Nunito", bold=True, size=10, color=NO_TX)))

n = fin + 2
ws.cell(row=n, column=1, value="Ratio máximo del conjunto").font = Font(name="Overpass", bold=True, size=10, color=NAVY)
ws.cell(row=n, column=5, value=f'=IF(COUNT(E{ini}:E{fin})=0,"",MAX(E{ini}:E{fin}))').number_format = "0.00"
ws.cell(row=n, column=5).font = Font(name="Overpass", bold=True, size=11, color=NAVY)
ws.cell(row=n, column=6, value=f'=IF(E{n}="","",IF(E{n}<=1,"CUMPLE","NO CUMPLE"))')
ws.cell(row=n, column=6).font = Font(name="Overpass", bold=True, size=10)
ws.cell(row=n, column=6).alignment = Alignment(horizontal="center")

nota = ("El ratio y el veredicto son aritmética, no ingeniería: comparan lo que tú escribiste. "
        "La solicitación, la resistencia y la norma que aplica las decides tú.")
ws.cell(row=n + 2, column=1, value=nota).font = Font(name="Nunito", size=9, color=GREY, italic=True)
ws.merge_cells(start_row=n + 2, start_column=1, end_row=n + 2, end_column=7)

# ── ELS ──────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("ELS · Servicio")
vac2 = [["[ elemento o nivel ]", "[ deflexión / deriva / vibración ]", None, None, None, None, "[ norma y artículo ]"] for _ in range(12)]
i2, f2 = hoja(ws2, "Resumen de verificaciones · Estado límite de servicio",
    "La sección que más se olvida, y la que el usuario del edificio nota primero.",
    ["Elemento / nivel", "Verificación", "Valor\nobtenido", "Límite\nadmisible", "Ratio\nobt/lím", "Veredicto", "Referencia normativa"],
    [22, 24, 13, 13, 10, 13, 30], vac2)

for r in range(i2, f2 + 1):
    ws2.cell(row=r, column=5, value=f'=IF(N(D{r})=0,"",C{r}/D{r})').number_format = "0.00"
    ws2.cell(row=r, column=6, value=f'=IF(E{r}="","",IF(E{r}<=1,"CUMPLE","NO CUMPLE"))')
    ws2.cell(row=r, column=6).font = Font(name="Overpass", bold=True, size=9)
    ws2.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="center")

r2 = f"F{i2}:F{f2}"
ws2.conditional_formatting.add(r2, CellIsRule(operator="equal", formula=['"CUMPLE"'],
    fill=PatternFill("solid", fgColor=OK_BG), font=Font(name="Overpass", bold=True, size=9, color=OK_TX)))
ws2.conditional_formatting.add(r2, CellIsRule(operator="equal", formula=['"NO CUMPLE"'],
    fill=PatternFill("solid", fgColor=NO_BG), font=Font(name="Overpass", bold=True, size=9, color=NO_TX)))

ws2.cell(row=f2 + 2, column=1,
    value="Una estructura puede cumplir todos los ELU y aun así ser inhabitable. Por eso esta hoja va aparte y no como un bloque más de la anterior."
    ).font = Font(name="Nunito", size=9, color=GREY, italic=True)
ws2.merge_cells(start_row=f2 + 2, start_column=1, end_row=f2 + 2, end_column=7)

# ── Instrucciones ────────────────────────────────────────────────────────
ws3 = wb.create_sheet("Cómo se usa")
ws3.sheet_view.showGridLines = False
ws3.column_dimensions["A"].width = 4
ws3.column_dimensions["B"].width = 96
filas = [
    ("t", "Cómo se usa esta hoja"),
    ("p", "Acompaña a la plantilla de memoria de cálculo. Rellena aquí el resumen y lleva la tabla ya resuelta a las secciones 10 y 11 del documento."),
    ("h", "Lo que hace"),
    ("l", "Calcula el ratio entre lo que escribiste como solicitación y lo que escribiste como resistencia."),
    ("l", "Marca CUMPLE cuando ese ratio es menor o igual a 1,00, y NO CUMPLE cuando lo supera."),
    ("l", "Te da el ratio máximo del conjunto, que es el número que gobierna."),
    ("h", "Lo que NO hace, a propósito"),
    ("l", "No propone secciones ni dimensiones."),
    ("l", "No conoce ninguna norma ni sus límites: el límite admisible lo escribes tú."),
    ("l", "No decide si tu modelo o tus hipótesis son correctos."),
    ("p", "En el momento en que la hoja propusiera un valor de diseño, dejaría de ser una plantilla documental y tendría que responder por ese número. La responsabilidad de lo que se firma no se delega en una hoja de cálculo."),
    ("h", "La columna que más se salta la gente"),
    ("p", "«Referencia normativa». Es la que convierte una tabla de números en una verificación que el revisor puede comprobar. Sin ella, la fila dice que algo cumple pero no contra qué."),
    ("d", "Plantilla de estructura documental de Design Modeling Academy. No es un método de cálculo ni una norma, y no sustituye el criterio ni la responsabilidad del ingeniero que firma."),
]
r = 2
for tipo, txt in filas:
    c = ws3.cell(row=r, column=2, value=("•  " + txt) if tipo == "l" else txt)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if tipo == "t":
        c.font = Font(name="Overpass", bold=True, size=16, color=NAVY); ws3.row_dimensions[r].height = 24
    elif tipo == "h":
        c.font = Font(name="Overpass", bold=True, size=10, color=ORANGE); r += 1
    elif tipo == "d":
        c.font = Font(name="Nunito", size=9, color=GREY, italic=True); ws3.row_dimensions[r].height = 40; r += 1
    else:
        c.font = Font(name="Nunito", size=11); ws3.row_dimensions[r].height = 34
    r += 1

OUT = "memoria-calculo/Resumen-Verificaciones-DMA.xlsx"
wb.save(OUT)
print("OK →", OUT, "·", len(wb.sheetnames), "hojas:", ", ".join(wb.sheetnames))
